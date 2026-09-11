from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import numpy as np
import polars as pl
from openpyxl import load_workbook

from valforecast.features.election_history import PARTIES
from valforecast.ingest.elections import read_val2022_district_results


def replace_zero_shares(
    shares: np.ndarray,
    *,
    valid_votes: int,
    pseudo_count: float = 0.5,
) -> np.ndarray:
    counts = shares * valid_votes
    counts = np.where(counts == 0, pseudo_count, counts)
    return np.asarray(counts / counts.sum(), dtype=float)


EXPECTED_RELATIONS = {
    "SAME": 4937,
    "COMPARABLE": 87,
    "MERGED": 35,
    "UNMATCHED": 1253,
}
EXPECTED_TARGET_DISTRICTS = 6312


@dataclass(frozen=True)
class ForecastUniverse:
    districts: pl.DataFrame
    previous_matrix: np.ndarray
    previous_valid_votes: np.ndarray
    eligible_weights: np.ndarray
    national_previous: np.ndarray
    national_valid_votes: int
    mapping_diagnostics: pl.DataFrame


def read_official_val2022_2026_crosswalk(path: Path) -> pl.DataFrame:
    sheet = load_workbook(path, read_only=True, data_only=True)["Jämförelser"]
    rows = sheet.iter_rows(values_only=True)
    next(rows)
    records: list[dict[str, object]] = []
    for values in rows:
        to_id = str(values[0]).strip()
        status = str(values[4] or "").strip()
        from_ids = [
            str(value).strip().zfill(8)
            for value in (values[5], values[6])
            if value is not None and str(value).strip()
        ]
        if status == "Kan jämföras":
            relation = "COMPARABLE" if from_ids and from_ids[0] != to_id else "SAME"
        elif status == "Kan jämföras mot flera":
            relation = "MERGED"
        else:
            relation = "UNMATCHED"
        records.append(
            {
                "to_district_id": to_id,
                "to_district_name": str(values[1] or "").strip(),
                "municipality_name": str(values[2] or "").strip(),
                "county_name": str(values[3] or "").strip(),
                "official_status": status,
                "relation": relation,
                "from_district_ids": from_ids,
            }
        )
    return pl.DataFrame(records)


def read_val2026_eligible_districts(path: Path) -> pl.DataFrame:
    sheet = load_workbook(path, read_only=True, data_only=True)["rostber_per_distrikt"]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() for value in next(rows)]
    index = {name: position for position, name in enumerate(headers)}
    records: list[dict[str, object]] = []
    for values in rows:
        district_id = values[index["Valdistriktskod"]]
        if district_id is None or str(values[index["Valdistrikt"]]).strip() == "Totalt":
            continue
        records.append(
            {
                "district_id": str(district_id).strip(),
                "district_name": str(values[index["Valdistrikt"]]).strip(),
                "municipality_id": str(values[index["Kommunkod"]]).strip().zfill(4),
                "municipality_name": str(values[index["Kommun"]]).strip(),
                "constituency_id": str(values[index["Riksdagsvalkretskod"]]).strip().zfill(2),
                "constituency_name": str(values[index["Riksdagsvalkrets"]]).strip(),
                "county_id": str(values[index["Kommunkod"]]).strip().zfill(4)[:2],
                "eligible_voters": int(str(values[index["Totalt"]])),
            }
        )
    frame = pl.DataFrame(records)
    if frame["district_id"].n_unique() != frame.height:
        raise ValueError("2026 eligible file has duplicate district identifiers")
    return frame


def _party_vector(shares: dict[str, float]) -> np.ndarray:
    return np.array([float(shares.get(party, 0.0)) for party in PARTIES], dtype=float)


def _district_share_map(
    results: pl.DataFrame,
) -> tuple[dict[str, tuple[np.ndarray, int]], dict[str, str]]:
    physical = results.filter(pl.col("district_kind") == "PHYSICAL")
    grouped = physical.group_by("district_id", "canonical_party_code").agg(
        pl.col("votes").sum().alias("votes"),
        pl.col("valid_votes").first().alias("valid_votes"),
        pl.col("municipality_id").first(),
    )
    by_district: dict[str, dict[str, float]] = {}
    valid: dict[str, int] = {}
    municipality: dict[str, str] = {}
    for record in grouped.iter_rows(named=True):
        district_id = str(record["district_id"])
        by_district.setdefault(district_id, {})
        by_district[district_id][str(record["canonical_party_code"])] = (
            float(record["votes"]) / float(record["valid_votes"])
            if record["valid_votes"]
            else 0.0
        )
        valid[district_id] = int(record["valid_votes"])
        municipality[district_id] = str(record["municipality_id"])
    return {
        district_id: (_party_vector(shares), valid[district_id])
        for district_id, shares in by_district.items()
    }, municipality


def unique_district_valid_votes(results: pl.DataFrame) -> pl.DataFrame:
    physical = results.filter(pl.col("district_kind") == "PHYSICAL")
    return physical.group_by("district_id").agg(
        pl.col("valid_votes").first().alias("valid_votes"),
        pl.col("municipality_id").first(),
    )


def national_valid_votes(results: pl.DataFrame) -> int:
    return int(unique_district_valid_votes(results).select(pl.col("valid_votes").sum()).item())


def missing_merged_predecessors(
    from_ids: list[str],
    available_district_ids: set[str],
) -> list[str]:
    if not from_ids:
        return ["<none listed>"]
    return [from_id for from_id in from_ids if from_id not in available_district_ids]


def _weighted_share_map(
    results: pl.DataFrame,
    group_column: str,
) -> dict[str, tuple[np.ndarray, int]]:
    physical = results.filter(pl.col("district_kind") == "PHYSICAL")
    grouped = physical.group_by(group_column, "canonical_party_code").agg(pl.col("votes").sum())
    district_valid = physical.group_by(group_column, "district_id").agg(
        pl.col("valid_votes").first().alias("valid_votes")
    )
    totals = district_valid.group_by(group_column).agg(
        pl.col("valid_votes").sum().alias("valid_votes")
    )
    valid = dict(totals.select(group_column, "valid_votes").iter_rows())
    shares: dict[str, dict[str, float]] = {}
    for record in grouped.iter_rows(named=True):
        key = str(record[group_column])
        shares.setdefault(key, {})
        shares[key][str(record["canonical_party_code"])] = (
            float(record["votes"]) / float(valid[key]) if valid[key] else 0.0
        )
    return {
        key: (_party_vector(party_shares), int(valid[key]))
        for key, party_shares in shares.items()
    }


def build_forecast_universe(root: Path) -> ForecastUniverse:
    mapping = read_official_val2022_2026_crosswalk(
        root / "data/raw/valmyndigheten/2026/valdistrikt_jamforelser_2022_2026.xlsx"
    )
    eligible = read_val2026_eligible_districts(
        root / "data/raw/valmyndigheten/2026/rostberattigade_riksdag_alder_kon_2026-08-14.xlsx"
    )
    results = read_val2022_district_results(
        root / "data/raw/valmyndigheten/2022/roster_per_distrikt_slutligt_riksdag.xlsx"
    )
    district_map, _district_municipality = _district_share_map(results)
    municipality_map = _weighted_share_map(results, "municipality_id")
    national_counts = results.filter(pl.col("district_kind") == "PHYSICAL").group_by(
        "canonical_party_code"
    ).agg(pl.col("votes").sum())
    national_valid = national_valid_votes(results)
    national_shares = {
        str(party): float(votes) / national_valid
        for party, votes in national_counts.select("canonical_party_code", "votes").iter_rows()
    }
    national_previous = _party_vector(national_shares)
    national_previous = national_previous / national_previous.sum()

    mapping_by_to = {row["to_district_id"]: row for row in mapping.iter_rows(named=True)}
    missing = set(eligible["district_id"]) - set(mapping_by_to)
    extra = set(mapping_by_to) - set(eligible["district_id"])
    if missing or extra:
        raise ValueError(
            f"2026 eligible/mapping mismatch: missing={sorted(missing)[:5]}"
        )

    rows: list[dict[str, object]] = []
    previous_rows: list[np.ndarray] = []
    previous_valid: list[int] = []
    weights: list[int] = []
    for record in eligible.sort("district_id").iter_rows(named=True):
        mapped = mapping_by_to[record["district_id"]]
        from_ids = list(mapped["from_district_ids"])
        relation = str(mapped["relation"])
        merge_missing: list[str] = []
        single = len(from_ids) == 1 and from_ids and from_ids[0] in district_map
        if relation in {"SAME", "COMPARABLE"} and single:
            vector, valid_votes = district_map[from_ids[0]]
            source = "official_comparable"
        elif relation == "MERGED":
            missing_preds = missing_merged_predecessors(from_ids, set(district_map))
            if missing_preds:
                merge_missing = missing_preds
                vector, valid_votes, source, _fallback_relation = _fallback(
                    record["municipality_id"], municipality_map, national_previous, national_valid
                )
                source = (
                    "municipality_2022_fallback_incomplete_merge"
                    if source == "municipality_2022_fallback"
                    else "national_2022_fallback_incomplete_merge"
                )
            else:
                parts = [district_map[from_id] for from_id in from_ids]
                totals = np.array([valid for _, valid in parts], dtype=float)
                vector = sum(share * valid for (share, valid) in parts) / totals.sum()
                valid_votes = int(totals.sum())
                source = "official_merge"
        else:
            vector, valid_votes, source, _fallback_relation = _fallback(
                record["municipality_id"], municipality_map, national_previous, national_valid
            )
        smoothed = replace_zero_shares(vector, valid_votes=max(valid_votes, 1))
        previous_rows.append(smoothed)
        previous_valid.append(valid_votes)
        weights.append(int(record["eligible_voters"]))
        rows.append(
            {
                **record,
                "relation": relation,
                "mapping_rule": source,
                "official_status": mapped["official_status"],
                "from_district_count": len(from_ids),
                "merge_missing_predecessors": ",".join(merge_missing),
                "previous_valid_votes": valid_votes,
            }
        )

    districts = pl.DataFrame(rows)
    _validate_forecast_universe(districts, np.array(weights, dtype=float), eligible, mapping)
    diagnostics = districts.group_by("official_status", "mapping_rule", "relation").len()
    return ForecastUniverse(
        districts=districts,
        previous_matrix=np.vstack(previous_rows),
        previous_valid_votes=np.array(previous_valid, dtype=int),
        eligible_weights=np.array(weights, dtype=float),
        national_previous=national_previous,
        national_valid_votes=national_valid,
        mapping_diagnostics=diagnostics.sort("official_status", "mapping_rule"),
    )


def _validate_forecast_universe(
    districts: pl.DataFrame,
    weights: np.ndarray,
    eligible: pl.DataFrame,
    mapping: pl.DataFrame,
) -> None:
    if districts.height != EXPECTED_TARGET_DISTRICTS:
        raise ValueError(f"Expected {EXPECTED_TARGET_DISTRICTS} target districts")
    if set(districts["district_id"]) != set(eligible["district_id"]):
        raise ValueError("Universe districts do not match the 2026 eligible file")
    extra = set(mapping["to_district_id"]) - set(eligible["district_id"])
    if extra:
        raise ValueError(f"Mapping has extra districts: {sorted(extra)[:5]}")
    if not np.all(weights > 0):
        raise ValueError("Eligible-voter weights must be strictly positive")
    counts = dict(districts.group_by("relation").len().iter_rows())
    if counts != EXPECTED_RELATIONS:
        raise ValueError(f"Unexpected relation counts: {counts}")
    allowed_rules = {
        "official_comparable",
        "official_merge",
        "municipality_2022_fallback",
        "national_2022_fallback",
        "municipality_2022_fallback_incomplete_merge",
        "national_2022_fallback_incomplete_merge",
    }
    unknown = set(districts["mapping_rule"].unique()) - allowed_rules
    if unknown:
        raise ValueError(f"Unknown mapping rules: {sorted(unknown)}")
    unmatched = districts.filter(pl.col("relation") == "UNMATCHED")
    fallback_rules = set(unmatched["mapping_rule"].unique())
    if not fallback_rules.issubset({"municipality_2022_fallback", "national_2022_fallback"}):
        raise ValueError(f"UNMATCHED districts have unexpected fallbacks: {fallback_rules}")


def _fallback(
    municipality_id: str,
    municipality_map: dict[str, tuple[np.ndarray, int]],
    national_previous: np.ndarray,
    national_valid: int,
) -> tuple[np.ndarray, int, str, str]:
    if municipality_id in municipality_map:
        vector, valid_votes = municipality_map[municipality_id]
        return vector, valid_votes, "municipality_2022_fallback", "UNMATCHED"
    return national_previous, national_valid, "national_2022_fallback", "UNMATCHED"
