from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

import polars as pl
from openpyxl import load_workbook

CANONICAL_PARTY_CODES = ("V", "S", "MP", "C", "L", "M", "KD", "SD", "OTHER")
PARTY_ALIASES = {
    "ARBETAREPARTIET-SOCIALDEMOKRATERNA": "S",
    "CENTERPARTIET": "C",
    "FP": "L",
    "FOLKPARTIET": "L",
    "FOLKPARTIET LIBERALERNA": "L",
    "KRISTDEMOKRATERNA": "KD",
    "LIBERALERNA": "L",
    "LIBERALERNA (TIDIGARE FOLKPARTIET)": "L",
    "MILJÖPARTIET DE GRÖNA": "MP",
    "MODERATERNA": "M",
    "SVERIGEDEMOKRATERNA": "SD",
    "VÄNSTERPARTIET": "V",
}

ELECTION_RESULT_COLUMNS = (
    "election_year",
    "election_date",
    "district_id",
    "district_version",
    "municipality_id",
    "county_id",
    "constituency_id",
    "party",
    "raw_party_code",
    "canonical_party_code",
    "votes",
    "valid_votes",
    "invalid_votes",
    "eligible_voters",
    "turnout",
    "vote_share",
)

LEGACY_ELECTION_DATES = {
    2010: date(2010, 9, 19),
    2014: date(2014, 9, 14),
}


def canonical_party_expression(column: str = "raw_party_code") -> pl.Expr:
    normalized = pl.col(column).cast(pl.String).str.strip_chars().str.to_uppercase()
    return (
        normalized.replace_strict(PARTY_ALIASES, default=normalized)
        .replace_strict(
            {code: code for code in CANONICAL_PARTY_CODES},
            default=pl.lit("OTHER"),
        )
        .alias("canonical_party_code")
    )


def add_derived_election_fields(frame: pl.DataFrame) -> pl.DataFrame:
    required = {
        "votes",
        "valid_votes",
        "invalid_votes",
        "eligible_voters",
        "raw_party_code",
    }
    missing = required - set(frame.columns)
    if missing:
        raise ValueError(f"Missing election columns: {sorted(missing)}")
    return frame.with_columns(
        canonical_party_expression(),
        pl.when(pl.col("valid_votes") > 0)
        .then(pl.col("votes") / pl.col("valid_votes"))
        .otherwise(None)
        .alias("vote_share"),
        pl.when(pl.col("eligible_voters") > 0)
        .then(
            (pl.col("valid_votes") + pl.col("invalid_votes").fill_null(0))
            / pl.col("eligible_voters")
        )
        .otherwise(None)
        .alias("turnout"),
    )


def election_quality_issues(frame: pl.DataFrame, tolerance: int = 1) -> pl.DataFrame:
    required = {
        "election_year",
        "district_id",
        "canonical_party_code",
        "votes",
        "valid_votes",
        "eligible_voters",
        "vote_share",
        "turnout",
    }
    missing = required - set(frame.columns)
    if missing:
        raise ValueError(f"Cannot validate election frame; missing: {sorted(missing)}")

    district = frame.group_by("election_year", "district_id").agg(
        pl.col("votes").sum().alias("party_votes"),
        pl.col("valid_votes").first().alias("valid_votes"),
        pl.col("eligible_voters").first().alias("eligible_voters"),
        pl.col("turnout").first().alias("turnout"),
        pl.col("vote_share").sum().alias("share_sum"),
        pl.col("raw_party_code").n_unique().alias("unique_parties"),
        pl.len().alias("party_rows"),
    )
    return district.with_columns(
        (pl.col("party_votes") - pl.col("valid_votes")).abs().alias("vote_difference"),
        ((pl.col("party_votes") - pl.col("valid_votes")).abs() > tolerance).alias(
            "party_votes_mismatch"
        ),
        (pl.col("eligible_voters") < pl.col("valid_votes")).alias("eligible_below_valid"),
        (~pl.col("turnout").is_between(0, 1, closed="both")).alias("invalid_turnout"),
        ((pl.col("share_sum") - 1).abs() > 1e-6).alias("invalid_share_sum"),
        (pl.col("unique_parties") != pl.col("party_rows")).alias("duplicate_party"),
    ).filter(
        pl.any_horizontal(
            "party_votes_mismatch",
            "eligible_below_valid",
            "invalid_turnout",
            "invalid_share_sum",
            "duplicate_party",
        )
    )


def read_val2018_district_results(path: Path) -> pl.DataFrame:
    sheet = load_workbook(path, read_only=True, data_only=True)["R antal"]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() for value in next(rows)]
    valid_votes_index = headers.index("RÖSTER GILTIGA")
    party_headers = headers[8 : headers.index("ÖVR") + 1]
    records: list[dict[str, object]] = []

    for values in rows:
        county_code = _excel_int(values[0])
        municipality_code = _excel_int(values[1])
        district_code = _excel_int(values[3])
        valid_votes = _excel_int(values[valid_votes_index])
        voters = _excel_int(values[headers.index("RÖSTANDE")])
        eligible = _excel_int(values[headers.index("RÖSTBERÄTTIGADE")])
        district_kind = "COLLECTION" if eligible == 0 else "PHYSICAL"
        district_id = district_id_2018(
            county_code,
            municipality_code,
            district_code,
            collection=district_kind == "COLLECTION",
        )
        for offset, raw_party_code in enumerate(party_headers, start=8):
            records.append(
                {
                    "election_year": 2018,
                    "election_date": date(2018, 9, 9),
                    "district_id": district_id,
                    "district_version": "VD_2018",
                    "district_kind": district_kind,
                    "municipality_id": f"{county_code:02d}{municipality_code:02d}",
                    "county_id": f"{county_code:02d}",
                    "constituency_id": f"{_excel_int(values[2]):02d}",
                    "district_name": str(values[7]).strip(),
                    "party": raw_party_code,
                    "raw_party_code": raw_party_code,
                    "votes": _excel_int(values[offset], none_as_zero=True),
                    "valid_votes": valid_votes,
                    "invalid_votes": voters - valid_votes,
                    "eligible_voters": eligible,
                }
            )
    return add_derived_election_fields(pl.DataFrame(records))


def read_legacy_district_results(path: Path, *, election_year: int) -> pl.DataFrame:
    """Read Valmyndigheten's semicolon district files for 2010 or 2014."""
    if election_year not in LEGACY_ELECTION_DATES:
        raise ValueError(f"Unsupported legacy election year: {election_year}")
    frame = pl.read_csv(
        path,
        separator=";",
        encoding="windows-1252",
        infer_schema_length=0,
    )
    name_column = "VALDISTRIKT" if election_year == 2010 else "Valdistrikt"
    party_columns = [
        column
        for column in frame.columns
        if column.endswith(" tal") and column not in {"BL tal", "OG tal"}
    ]
    records: list[dict[str, object]] = []
    for row in frame.iter_rows(named=True):
        county = str(row["LAN"]).zfill(2)
        municipality = str(row["KOM"]).zfill(2)
        eligible = int(str(row["Rostb"] or 0))
        raw_district = str(row["VALDIST"])
        district = (
            raw_district.zfill(4) if eligible > 0 else raw_district.zfill(2)
        )
        valid_votes = int(str(row["Rost Giltiga"]))
        votes_cast = int(str(row["Rostande"]))
        district_id = f"{county}{municipality}{district}"
        for party_column in party_columns:
            raw_party_code = party_column.removesuffix(" tal")
            records.append(
                {
                    "election_year": election_year,
                    "election_date": LEGACY_ELECTION_DATES[election_year],
                    "district_id": district_id,
                    "district_version": f"VD_{election_year}",
                    "district_kind": "COLLECTION" if eligible == 0 else "PHYSICAL",
                    "municipality_id": f"{county}{municipality}",
                    "county_id": county,
                    "constituency_id": str(row["RIKSDAGSVALKRETS"]).strip(),
                    "district_name": str(row[name_column]).strip(),
                    "party": raw_party_code,
                    "raw_party_code": raw_party_code,
                    "votes": int(str(row[party_column] or 0)),
                    "valid_votes": valid_votes,
                    "invalid_votes": votes_cast - valid_votes,
                    "eligible_voters": eligible,
                }
            )
    return add_derived_election_fields(pl.DataFrame(records))


def read_val2022_district_results(path: Path) -> pl.DataFrame:
    sheet = load_workbook(path, read_only=True, data_only=True)["roster_RD"]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() for value in next(rows)]
    index = {name: position for position, name in enumerate(headers)}
    raw_records = [tuple(values) for values in rows]

    totals: dict[str, dict[str, int]] = {}
    for values in raw_records:
        district = str(values[index["Distrikt"]]).strip()
        label = str(values[index["Parti"]]).strip()
        district_totals = totals.setdefault(
            district,
            {
                "valid_votes": 0,
                "invalid_votes": 0,
                "eligible_voters": _excel_int(values[index["Röstberättigade"]], none_as_zero=True),
            },
        )
        votes = _excel_int(values[index["Röster"]], none_as_zero=True)
        if label == "Summa giltiga röster":
            district_totals["valid_votes"] = votes
        elif label in {"blanka röster", "övriga ogiltiga", "ej anmält deltagande"}:
            district_totals["invalid_votes"] += votes

    excluded = {
        "Summa giltiga röster",
        "Valdeltagande",
        "blanka röster",
        "övriga ogiltiga",
        "ej anmält deltagande",
    }
    records: list[dict[str, object]] = []
    for values in raw_records:
        raw_party = str(values[index["Parti"]]).strip()
        if raw_party in excluded:
            continue
        raw_district = str(values[index["Distrikt"]]).strip()
        district_id = str(values[index["Valdistriktskod"]]).strip()
        district_totals = totals[raw_district]
        district_kind = "COLLECTION" if district_totals["eligible_voters"] == 0 else "PHYSICAL"
        records.append(
            {
                "election_year": 2022,
                "election_date": date(2022, 9, 11),
                "district_id": district_id,
                "district_version": "VD_2022",
                "district_kind": district_kind,
                "municipality_id": district_id[:4],
                "county_id": district_id[:2],
                "constituency_id": str(values[index["Valkretskod"]]).strip(),
                "district_name": str(values[index["Valdistriktnamn"]]).strip(),
                "party": raw_party,
                "raw_party_code": raw_party,
                "votes": _excel_int(values[index["Röster"]], none_as_zero=True),
                **district_totals,
            }
        )
    return add_derived_election_fields(pl.DataFrame(records))


def _excel_int(value: object, *, none_as_zero: bool = False) -> int:
    if value is None and none_as_zero:
        return 0
    if isinstance(value, bool):
        raise ValueError("Boolean value found where integer was expected")
    if isinstance(value, (int, float, Decimal, str)):
        return int(value)
    raise ValueError(f"Unsupported Excel integer value: {value!r}")


def district_id_2018(
    county_code: int,
    municipality_code: int,
    district_code: int,
    *,
    collection: bool,
) -> str:
    district_width = 2 if collection else 4
    return f"{county_code:02d}{municipality_code:02d}{district_code:0{district_width}d}"
