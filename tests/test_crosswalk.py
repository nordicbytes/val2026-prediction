from pathlib import Path

import polars as pl
from openpyxl import Workbook

from valforecast.geo.crosswalk import (
    classify_official_pairs,
    directly_comparable,
    read_official_val2018_2022_crosswalk,
)


def test_classifies_graph_relations() -> None:
    pairs = pl.DataFrame(
        {
            "from_district_id": ["A", "B", "B", "C", "D"],
            "to_district_id": ["A", "X", "Y", "Z", "Z"],
        }
    )

    result = classify_official_pairs(pairs, from_election=2018, to_election=2022)
    relations = {
        (row["from_district_id"], row["to_district_id"]): row["relation"]
        for row in result.iter_rows(named=True)
    }

    assert relations == {
        ("A", "A"): "SAME",
        ("B", "X"): "SPLIT",
        ("B", "Y"): "SPLIT",
        ("C", "Z"): "MERGED",
        ("D", "Z"): "MERGED",
    }
    assert directly_comparable(result)["from_district_id"].to_list() == ["A"]


def test_renamed_one_to_one_district_is_comparable() -> None:
    pairs = pl.DataFrame({"from_district_id": ["old"], "to_district_id": ["new"]})
    result = classify_official_pairs(pairs, from_election=2018, to_election=2022)
    assert result.item(0, "relation") == "COMPARABLE"


def test_reads_all_official_comparability_forms(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Fysiska valdistrikt"
    sheet.append(
        [
            "Kod_2022",
            "Valdistrikt",
            "Jämförbart",
            "Valdistriktskod2018",
            "Kod 2 2018",
            "Kod 3 2018",
            "Valdistriktsnamn 2018",
            "Namn 2 2018",
            "Namn 3 2018",
            "Län",
        ]
    )
    sheet.append(["01140101", "A", "ja", "01140101", None, None, "A", None, None, "X"])
    sheet.append(["01170009", "B", "01170313", "01170313", None, None, "B", None, None, "X"])
    sheet.append(
        [
            "06621621",
            "C",
            "06620511, 06620812",
            "06620511",
            "06620812",
            None,
            "C1",
            "C2",
            None,
            "X",
        ]
    )
    sheet.append(["01140231", "D", "nej", None, None, None, None, None, None, "X"])
    path = tmp_path / "mapping.xlsx"
    workbook.save(path)

    result = read_official_val2018_2022_crosswalk(path)
    counts = dict(result.group_by("relation").len().iter_rows())

    assert counts == {"SAME": 1, "COMPARABLE": 1, "MERGED": 2, "NOT_COMPARABLE": 1}
    assert result["to_district_id"].n_unique() == 4

